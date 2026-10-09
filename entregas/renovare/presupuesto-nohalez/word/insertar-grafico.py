"""Inserta el gráfico de barras nativo de Word «Peso de cada capítulo sobre el total».

Uso: python3 insertar-grafico.py ["ruta.docx"]
Sustituye el párrafo marcador §GRAFICO_CAPITULOS§ por un gráfico de Word con su hoja de
datos incrustada. En Word: clic derecho sobre el gráfico › Editar datos. La hoja calcula el
% de cada capítulo a partir de su importe, así que al cambiar un importe la barra se ajusta sola.
"""
import io, os, re, sys, zipfile
from xml.sax.saxutils import escape
from openpyxl import Workbook
from openpyxl.styles import Font

AQUI = os.path.dirname(os.path.abspath(__file__))
DOCX = sys.argv[1] if len(sys.argv) > 1 else os.path.join(AQUI, '..', 'PRESUPUESTO NOHALEZ - SV-0045-03.docx')
MARCADOR = '§GRAFICO_CAPITULOS§'

CAPITULOS = [  # (capítulo, importe sin IVA)
    ('Suelos: nivelación y microcemento', 11211),
    ('Electricidad, fontanería y ayudas', 9845),
    ('Paredes y techos', 7854),
    ('Baño completo', 5897),
    ('Cocina: mobiliario y encimera', 5114),
    ('Puertas y ventanas', 2902),
    ('Climatización', 2577),
    ('Protección, limpieza y residuos', 959),
]
ORO, GRIS, TINTA, LINEA = 'A3843F', '625B50', '1C1A16', 'E2DBCD'
ANCHO_MM, ALTO_MM = 170, 56
EMU = 36000


def hoja_datos() -> bytes:
    wb = Workbook(); ws = wb.active; ws.title = 'Datos'
    ws.append(['Capítulo', 'Importe sin IVA (€)', '% del total'])
    for c in ws[1]:
        c.font = Font(bold=True)
    n = len(CAPITULOS)
    for i, (cap, imp) in enumerate(CAPITULOS, start=2):
        ws.append([cap, imp, f'=B{i}/$B${n + 2}'])
        ws[f'B{i}'].number_format = '#,##0 €'; ws[f'C{i}'].number_format = '0.0%'
    ws.append(['Total proyecto base', f'=SUM(B2:B{n + 1})', f'=SUM(C2:C{n + 1})'])
    ws[f'B{n + 2}'].number_format = '#,##0 €'; ws[f'C{n + 2}'].number_format = '0.0%'
    ws.column_dimensions['A'].width = 38; ws.column_dimensions['B'].width = 20; ws.column_dimensions['C'].width = 14
    bio = io.BytesIO(); wb.save(bio); return bio.getvalue()


def txpr(sz, color, bold=False):
    b = ' b="1"' if bold else ''
    return (f'<c:txPr><a:bodyPr/><a:lstStyle/><a:p><a:pPr><a:defRPr sz="{sz}"{b}><a:solidFill><a:srgbClr val="{color}"/></a:solidFill>'
            f'<a:latin typeface="Inter"/><a:cs typeface="Inter"/></a:defRPr></a:pPr><a:endParaRPr lang="es-ES"/></a:p></c:txPr>')


def grafico_xml() -> str:
    n = len(CAPITULOS); total = sum(v for _, v in CAPITULOS)
    cats = ''.join(f'<c:pt idx="{i}"><c:v>{escape(c)}</c:v></c:pt>' for i, (c, _) in enumerate(CAPITULOS))
    vals = ''.join(f'<c:pt idx="{i}"><c:v>{v / total:.6f}</c:v></c:pt>' for i, (_, v) in enumerate(CAPITULOS))
    sin = '<c:spPr><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>'
    return f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
<c:date1904 val="0"/><c:lang val="es-ES"/><c:roundedCorners val="0"/>
<c:chart><c:autoTitleDeleted val="1"/><c:plotArea><c:layout/>
<c:barChart><c:barDir val="bar"/><c:grouping val="clustered"/><c:varyColors val="0"/>
<c:ser><c:idx val="0"/><c:order val="0"/>
<c:tx><c:strRef><c:f>Datos!$C$1</c:f><c:strCache><c:ptCount val="1"/><c:pt idx="0"><c:v>% del total</c:v></c:pt></c:strCache></c:strRef></c:tx>
<c:spPr><a:solidFill><a:srgbClr val="{ORO}"/></a:solidFill><a:ln><a:noFill/></a:ln></c:spPr><c:invertIfNegative val="0"/>
<c:dLbls><c:numFmt formatCode="0.0%" sourceLinked="0"/>{sin}{txpr(800, TINTA, True)}<c:dLblPos val="outEnd"/>
<c:showLegendKey val="0"/><c:showVal val="1"/><c:showCatName val="0"/><c:showSerName val="0"/><c:showPercent val="0"/><c:showBubbleSize val="0"/></c:dLbls>
<c:cat><c:strRef><c:f>Datos!$A$2:$A${n + 1}</c:f><c:strCache><c:ptCount val="{n}"/>{cats}</c:strCache></c:strRef></c:cat>
<c:val><c:numRef><c:f>Datos!$C$2:$C${n + 1}</c:f><c:numCache><c:formatCode>0.0%</c:formatCode><c:ptCount val="{n}"/>{vals}</c:numCache></c:numRef></c:val>
</c:ser><c:gapWidth val="80"/><c:axId val="501"/><c:axId val="502"/></c:barChart>
<c:catAx><c:axId val="501"/><c:scaling><c:orientation val="maxMin"/></c:scaling><c:delete val="0"/><c:axPos val="l"/>
<c:numFmt formatCode="General" sourceLinked="1"/><c:majorTickMark val="none"/><c:minorTickMark val="none"/><c:tickLblPos val="nextTo"/>
<c:spPr><a:ln w="6350"><a:solidFill><a:srgbClr val="{LINEA}"/></a:solidFill></a:ln></c:spPr>{txpr(800, GRIS)}
<c:crossAx val="502"/><c:crosses val="autoZero"/><c:auto val="1"/><c:lblAlgn val="ctr"/><c:lblOffset val="100"/><c:noMultiLvlLbl val="0"/></c:catAx>
<c:valAx><c:axId val="502"/><c:scaling><c:orientation val="minMax"/><c:min val="0"/></c:scaling><c:delete val="1"/><c:axPos val="t"/>
<c:numFmt formatCode="0%" sourceLinked="0"/><c:majorTickMark val="none"/><c:minorTickMark val="none"/><c:tickLblPos val="nextTo"/>
<c:crossAx val="501"/><c:crosses val="autoZero"/><c:crossBetween val="between"/></c:valAx>
{sin}</c:plotArea><c:plotVisOnly val="1"/><c:dispBlanksAs val="gap"/></c:chart>
{sin}{txpr(800, GRIS)}
<c:externalData r:id="rId1"><c:autoUpdate val="0"/></c:externalData>
</c:chartSpace>'''


def main():
    with zipfile.ZipFile(DOCX) as z:
        partes = {n: z.read(n) for n in z.namelist()}
    doc = partes['word/document.xml'].decode()
    m = None
    for pm in re.finditer(r'<w:p\b(?:(?!<w:p\b).)*?</w:p>', doc, re.S):
        if MARCADOR in pm.group(0):
            m = pm; break
    if not m:
        print('Marcador no encontrado (¿gráfico ya insertado?)'); return
    ppr = re.search(r'<w:pPr>.*?</w:pPr>', m.group(0), re.S)
    dibujo = (f'<w:p>{ppr.group(0) if ppr else ""}<w:r><w:drawing><wp:inline distT="0" distB="0" distL="0" distR="0">'
              f'<wp:extent cx="{ANCHO_MM * EMU}" cy="{ALTO_MM * EMU}"/><wp:effectExtent l="0" t="0" r="0" b="0"/>'
              '<wp:docPr id="9001" name="Gráfico de capítulos" descr="Peso de cada capítulo sobre el total del proyecto base"/>'
              '<wp:cNvGraphicFramePr/><a:graphic xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
              '<a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/chart">'
              '<c:chart xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
              'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="rIdGrafico1"/>'
              '</a:graphicData></a:graphic></wp:inline></w:drawing></w:r></w:p>')
    doc = doc[:m.start()] + dibujo + doc[m.end():]
    if 'xmlns:wp=' not in doc[:3000]:
        doc = doc.replace('<w:document ', '<w:document xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" ', 1)
    partes['word/document.xml'] = doc.encode()

    partes['word/charts/chart1.xml'] = grafico_xml().encode()
    partes['word/charts/_rels/chart1.xml.rels'] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/package" '
        'Target="../embeddings/Datos_grafico_capitulos.xlsx"/></Relationships>').encode()
    partes['word/embeddings/Datos_grafico_capitulos.xlsx'] = hoja_datos()

    rels = partes['word/_rels/document.xml.rels'].decode()
    rels = rels.replace('</Relationships>', '<Relationship Id="rIdGrafico1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/chart" Target="charts/chart1.xml"/></Relationships>')
    partes['word/_rels/document.xml.rels'] = rels.encode()
    ct = partes['[Content_Types].xml'].decode()
    if 'Extension="xlsx"' not in ct:
        ct = ct.replace('<Default ', '<Default Extension="xlsx" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"/><Default ', 1)
    ct = ct.replace('</Types>', '<Override PartName="/word/charts/chart1.xml" ContentType="application/vnd.openxmlformats-officedocument.drawingml.chart+xml"/></Types>')
    partes['[Content_Types].xml'] = ct.encode()

    tmp = DOCX + '.tmp'
    with zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as z:
        z.writestr('[Content_Types].xml', partes.pop('[Content_Types].xml'))
        for n, d in partes.items():
            z.writestr(n, d)
    os.replace(tmp, DOCX)
    print('Gráfico insertado en', DOCX)


if __name__ == '__main__':
    main()
